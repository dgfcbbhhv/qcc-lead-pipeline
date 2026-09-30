# -*- coding: utf-8 -*-
import os as _os
HERE = _os.path.dirname(_os.path.abspath(__file__))
BASE = _os.environ.get("LEAD_HOME") or HERE
"""手机号数据清洗：合并批次 -> 号码分流 -> 剔除无效 -> 标记跨企业重复号码"""
import json, os, re
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def load(p):
    fp = os.path.join(BASE, p)
    if not os.path.exists(fp):
        return {}
    with open(fp, 'r', encoding='utf-8') as f:
        return json.load(f)

import glob
data = {}
data.update(load('qcc_results_all.json'))     # 早期单批结果（可选，不存在就跳过）
# 邮箱新增批次：自动加载所有 qcc_results_batchN.json
batch_files = sorted(glob.glob(os.path.join(BASE, 'qcc_results_batch*.json')))
for bf in batch_files:
    data.update(load(os.path.basename(bf)))
print('批次文件:', [os.path.basename(b) for b in batch_files])
if not data:
    raise SystemExit('未找到任何 qcc_results_*.json，请先完成 MCP 查询步骤')

MOBILE_RE = re.compile(r'^1\d{10}$')

def split_numbers(rec):
    """把一条记录拆成：有效手机号 / 代记账手机号 / 座机 / 无效号码"""
    note = rec.get('note', '') or ''
    # 从 note 里解析出被标记为无效 / 疑似代记账的号码
    invalid_set = set(re.findall(r'(1\d{10}|\d{3,4}-\d{7,8})[^;；,，]{0,6}?标记无效', note))
    daiji_set = set(re.findall(r'(1\d{10})[^;；,，]{0,6}?疑似代记账', note))
    # note 里"两个均疑似代记账""后两个疑似代记账"这类描述，退化为整条标记
    if re.search(r'(均|全部|都)疑似代记账', note) or note.strip() == '疑似代记账':
        daiji_set |= {m for m in rec.get('mobile', []) if MOBILE_RE.match(m)}

    good, daiji, bad, land = [], [], [], []
    for m in rec.get('mobile', []):
        if not MOBILE_RE.match(m):
            land.append(m)
        elif m in invalid_set:
            bad.append(m)
        elif m in daiji_set:
            daiji.append(m)
        else:
            good.append(m)
    for l in rec.get('landline', []):
        (bad if l in invalid_set else land).append(l)
    # 去重保序
    uniq = lambda xs: list(dict.fromkeys(xs))
    return uniq(good), uniq(daiji), uniq(land), uniq(bad)

# ---------- 1. 逐条清洗 ----------
clean_rows = []
num_owners = defaultdict(set)   # 号码 -> 涉及企业集合
num_kind = {}                   # 号码 -> 首次出现时的类型

for name, rec in data.items():
    good, daiji, land, bad = split_numbers(rec)
    clean_rows.append({
        'name': name, 'good': good, 'daiji': daiji,
        'land': land, 'bad': bad, 'note': rec.get('note', '')
    })
    # 共用号码统计要覆盖有效号 + 代记账号（代账公司号码正是重复出现的那一类）
    for m in good:
        num_owners[m].add(name)
        num_kind.setdefault(m, '有效号')
    for m in daiji:
        num_owners[m].add(name)
        num_kind.setdefault(m, '疑似代记账')

clean_rows.sort(key=lambda r: r['name'])

# ---------- 2. 跨企业重复号码（代账公司特征） ----------
dups = {n: owners for n, owners in num_owners.items() if len(owners) >= 2}
dup_numbers = set(dups.keys())

# ---------- 3. 输出 Excel ----------
wb = Workbook()
thin = Side(style='thin', color='D0D0D0')
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
HF = PatternFill('solid', fgColor='2F5597')

def style_header(ws, cols):
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HF
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = bd
    ws.freeze_panes = 'A2'

# Sheet1 清洗总表
ws = wb.active
ws.title = '清洗总表'
style_header(ws, ['序号', '企业名称', '手机号1', '手机号2', '更多手机号',
                  '备用号(疑似代记账)', '座机', '已剔除无效号', '号码可信度', '备注'])

FILL = {
    '高': PatternFill('solid', fgColor='C6EFCE'),
    '中': PatternFill('solid', fgColor='FFF2CC'),
    '低': PatternFill('solid', fgColor='FCE4E4'),
}
stat = defaultdict(int)

for i, r in enumerate(clean_rows, 1):
    shared = [m for m in r['good'] if m in dup_numbers]
    private = [m for m in r['good'] if m not in dup_numbers]
    if len(r['good']) >= 2 and private:
        level = '高'
    elif len(private) == 1 or (len(r['good']) >= 2 and not private):
        level = '中'
    elif r['daiji'] or r['land']:
        level = '低'
    else:
        level = '低'
    stat[level] += 1

    extra_flags = []
    if shared:
        extra_flags.append(f'有 {len(shared)} 个号与其他企业共用')
    ws.append([
        i, r['name'],
        private[0] if private else (r['good'][0] if r['good'] else ''),
        private[1] if len(private) > 1 else '',
        '、'.join(private[2:]),
        '、'.join(r['daiji']),
        '、'.join(r['land']),
        '、'.join(r['bad']),
        level,
        ('; '.join(filter(None, [r['note'], *extra_flags]))).strip('; ')
    ])
    ws.cell(row=ws.max_row, column=9).fill = FILL[level]

for w, c in zip([6, 44, 15, 15, 26, 26, 30, 22, 11, 46], range(1, 11)):
    ws.column_dimensions[get_column_letter(c)].width = w
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10):
    for cell in row:
        cell.border = bd
        cell.alignment = Alignment(vertical='center', wrap_text=True) if cell.column in (2, 5, 6, 7, 8, 10) \
            else Alignment(horizontal='center', vertical='center')

# Sheet2 共用号码清单
ws2 = wb.create_sheet('共用号码(疑似代账)')
style_header(ws2, ['序号', '手机号', '号码类型', '涉及企业数', '涉及企业'])
for i, (num, owners) in enumerate(sorted(dups.items(), key=lambda kv: -len(kv[1])), 1):
    ws2.append([i, num, num_kind.get(num, ''), len(owners), ' / '.join(sorted(owners))])
for w, c in zip([6, 16, 14, 12, 110], range(1, 6)):
    ws2.column_dimensions[get_column_letter(c)].width = w
for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=5):
    for cell in row:
        cell.border = bd
        cell.alignment = Alignment(vertical='center', wrap_text=True)

# Sheet3 清洗统计
ws3 = wb.create_sheet('清洗统计')
style_header(ws3, ['项目', '数值'])
total = len(clean_rows)
n_good = sum(1 for r in clean_rows if r['good'])
n_daiji_only = sum(1 for r in clean_rows if not r['good'] and r['daiji'])
n_land_only = sum(1 for r in clean_rows if not r['good'] and not r['daiji'] and r['land'])
n_none = sum(1 for r in clean_rows if not r['good'] and not r['daiji'] and not r['land'])
all_mob = [m for r in clean_rows for m in r['good']]
n_removed = sum(len(r['bad']) for r in clean_rows)

rows3 = [
    ('已查企业总数', total),
    ('有有效手机号的企业', n_good),
    ('  其中 ≥2 个有效手机号', sum(1 for r in clean_rows if len(r['good']) >= 2)),
    ('  其中 仅1个有效手机号', sum(1 for r in clean_rows if len(r['good']) == 1)),
    ('只有代记账号码的企业', n_daiji_only),
    ('只有座机的企业', n_land_only),
    ('完全无联系方式的企业', n_none),
    ('有效手机号总数', len(all_mob)),
    ('去重后不重复手机号数', len(set(all_mob))),
    ('被剔除的无效号码数', n_removed),
    ('跨企业共用号码数(疑似代账号)', len(dups)),
    ('最多共用家数', max((len(v) for v in dups.values()), default=0)),
]
for k, v in rows3:
    ws3.append([k, v])
ws3.column_dimensions['A'].width = 34
ws3.column_dimensions['B'].width = 14
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=2):
    for cell in row:
        cell.border = bd
        cell.alignment = Alignment(horizontal='left' if cell.column == 1 else 'center', vertical='center')

out = os.path.join(BASE, '手机号清洗结果.xlsx')
wb.save(out)

# 导出清洗后的干净 JSON 供后续批次追加
clean_json = {r['name']: {'good': r['good'], 'daiji': r['daiji'],
                          'land': r['land'], 'bad': r['bad'], 'note': r['note']}
              for r in clean_rows}
with open(os.path.join(BASE, 'clean_results.json'), 'w', encoding='utf-8') as f:
    json.dump(clean_json, f, ensure_ascii=False, indent=1)

print('saved:', out)
for k, v in rows3:
    print(f'  {k}: {v}')
print('\n共用号码 TOP10:')
for num, owners in sorted(dups.items(), key=lambda kv: -len(kv[1]))[:10]:
    print(f'  {num}  涉及 {len(owners)} 家: {", ".join(list(sorted(owners))[:3])}...')
